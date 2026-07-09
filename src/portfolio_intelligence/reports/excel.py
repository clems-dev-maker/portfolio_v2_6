from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference, LineChart, BarChart
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from portfolio_intelligence.analytics import build_analytics
from portfolio_intelligence.risk.engine import build_risk
from portfolio_intelligence.performance.metrics import DashboardMetrics, monthly_flows
from portfolio_intelligence.performance.dashboard import build_performance
from portfolio_intelligence.lookthrough import build_lookthrough

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=18, bold=True, color="111827")
SUBTITLE_FONT = Font(size=11, color="6B7280")
THIN = Side(style="thin", color="E5E7EB")


def _format_sheet(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center")
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 48)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> tuple[int, int]:
    for j, col in enumerate(df.columns, start=start_col):
        cell = ws.cell(start_row, j, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for i, row in enumerate(df.astype(object).where(pd.notna(df), None).itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=start_col):
            ws.cell(i, j, value)
    return start_row, start_row + len(df)


def _add_pie(ws, title: str, labels_col: int, values_col: int, min_row: int, max_row: int, anchor: str) -> None:
    if max_row <= min_row:
        return
    chart = PieChart()
    chart.title = title
    data = Reference(ws, min_col=values_col, min_row=min_row, max_row=max_row)
    labels = Reference(ws, min_col=labels_col, min_row=min_row + 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    ws.add_chart(chart, anchor)


def _add_bar(ws, title: str, labels_col: int, values_col: int, min_row: int, max_row: int, anchor: str) -> None:
    if max_row <= min_row:
        return
    chart = BarChart()
    chart.type = "bar"
    chart.title = title
    data = Reference(ws, min_col=values_col, min_row=min_row, max_row=max_row)
    labels = Reference(ws, min_col=labels_col, min_row=min_row + 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    ws.add_chart(chart, anchor)


def write_report(output: str | Path, transactions: pd.DataFrame, tx_summary, net_worth, metrics: DashboardMetrics) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    analytics = build_analytics(transactions, net_worth)
    performance = build_performance(transactions, net_worth, metrics)
    lookthrough = build_lookthrough(transactions, net_worth)

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = "Portfolio Intelligence v2.6 — Reports Dashboard"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Rapport analytique généré depuis les imports Trade Republic"
    ws["A2"].font = SUBTITLE_FONT

    rows = [
        ("Valeur portefeuille", metrics.portfolio_value),
        ("Capital investi estimé", metrics.invested_estimate),
        ("Plus-value estimée", metrics.pnl_estimate),
        ("Plus-value estimée %", metrics.pnl_percent / 100),
        ("Transactions", metrics.transaction_count),
        ("Positions estimées", metrics.position_count_estimate),
        ("Première date", metrics.first_date),
        ("Dernière date", metrics.last_date),
        ("Frais", metrics.fees),
        ("Dividendes", metrics.dividends),
        ("Intérêts", metrics.interests),
    ]
    ws.append([])
    ws.append(["Indicateur", "Valeur"])
    for c in ws[4]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for r in rows:
        ws.append(list(r))
    for row in range(5, 5 + len(rows)):
        if isinstance(ws.cell(row, 2).value, (float, int)):
            ws.cell(row, 2).number_format = '#,##0.00'
    ws[8][1].number_format = '0.00%'
    _format_sheet(ws)

    # Transactions brutes
    tx_ws = wb.create_sheet("Transactions")
    _write_df(tx_ws, transactions)
    _format_sheet(tx_ws)

    # Positions estimées
    pos_ws = wb.create_sheet("Positions estimées")
    positions = analytics.positions_estimated.copy()
    _write_df(pos_ws, positions)
    for row in range(2, len(positions) + 2):
        for col in range(3, pos_ws.max_column + 1):
            if isinstance(pos_ws.cell(row, col).value, (float, int)):
                pos_ws.cell(row, col).number_format = '#,##0.00'
    _format_sheet(pos_ws)

    # Allocations
    alloc_ws = wb.create_sheet("Allocations")
    alloc_ws["A1"] = "Allocation par classe d'actifs"
    alloc_ws["A1"].font = TITLE_FONT
    r1, r2 = _write_df(alloc_ws, analytics.class_allocation, start_row=3)
    _add_pie(alloc_ws, "Classes d'actifs", 1, 2, r1, r2, "E3")
    alloc_ws["A15"] = "Allocation sectorielle"
    alloc_ws["A15"].font = TITLE_FONT
    r1, r2 = _write_df(alloc_ws, analytics.sector_allocation, start_row=17)
    _add_bar(alloc_ws, "Secteurs", 1, 2, r1, r2, "E17")
    alloc_ws["A32"] = "Allocation géographique"
    alloc_ws["A32"].font = TITLE_FONT
    r1, r2 = _write_df(alloc_ws, analytics.geography_allocation, start_row=34)
    _add_pie(alloc_ws, "Zones géographiques", 1, 2, r1, r2, "E34")
    alloc_ws["A47"] = "Allocation devises"
    alloc_ws["A47"].font = TITLE_FONT
    _write_df(alloc_ws, analytics.currency_allocation, start_row=49)
    for row in alloc_ws.iter_rows():
        for cell in row:
            if cell.column == 3 and isinstance(cell.value, (float, int)):
                cell.number_format = '0.00%'
    _format_sheet(alloc_ws)

    # Suivi mensuel analytique
    monthly = analytics.monthly_evolution
    month_ws = wb.create_sheet("Evolution mensuelle")
    _write_df(month_ws, monthly)
    for row in range(2, len(monthly) + 2):
        for col in range(2, month_ws.max_column + 1):
            if isinstance(month_ws.cell(row, col).value, (float, int)):
                month_ws.cell(row, col).number_format = '#,##0.00'
    if len(monthly) > 1:
        chart = LineChart()
        chart.title = "Capital investi vs valeur estimée"
        data = Reference(month_ws, min_col=3, max_col=4, min_row=1, max_row=len(monthly) + 1)
        cats = Reference(month_ws, min_col=1, min_row=2, max_row=len(monthly) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        month_ws.add_chart(chart, "G2")
    _format_sheet(month_ws)

    # Flux mensuels conservés pour compatibilité v2.1
    flows = monthly_flows(transactions)
    fl_ws = wb.create_sheet("Flux mensuels")
    _write_df(fl_ws, flows)
    if len(flows) > 1:
        chart = LineChart()
        chart.title = "Flux mensuels"
        data = Reference(fl_ws, min_col=2, min_row=1, max_row=len(flows) + 1)
        cats = Reference(fl_ws, min_col=1, min_row=2, max_row=len(flows) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        fl_ws.add_chart(chart, "D2")
    _format_sheet(fl_ws)


    # Performance Dashboard v2.4
    perf_ws = wb.create_sheet("Performance Dashboard")
    perf_ws["A1"] = "Performance Dashboard v2.4"
    perf_ws["A1"].font = TITLE_FONT
    perf_ws["A2"] = "Performance depuis l'origine : total return, CAGR, XIRR, meilleurs/pire mois et benchmark théorique."
    perf_ws["A2"].font = SUBTITLE_FONT
    _write_df(perf_ws, performance.summary, start_row=4)
    for row in range(5, 5 + len(performance.summary)):
        label = str(perf_ws.cell(row, 1).value or "")
        if isinstance(perf_ws.cell(row, 2).value, (float, int)):
            if "Capital" in label or "Valeur" in label:
                perf_ws.cell(row, 2).number_format = '#,##0.00 €'
            else:
                perf_ws.cell(row, 2).number_format = '0.00%'

    perf_ws["D4"] = "Rendements annuels"
    perf_ws["D4"].font = TITLE_FONT
    _write_df(perf_ws, performance.annual_returns, start_row=6, start_col=4)
    for row in range(7, 7 + len(performance.annual_returns)):
        perf_ws.cell(row, 5).number_format = '0.00%'
        perf_ws.cell(row, 6).number_format = '#,##0.00 €'
        perf_ws.cell(row, 7).number_format = '#,##0.00 €'
    if len(performance.annual_returns) > 0:
        _add_bar(perf_ws, "Rendements annuels", 4, 5, 6, 6 + len(performance.annual_returns), "I6")
    _format_sheet(perf_ws)

    pm_ws = wb.create_sheet("Rendements mensuels")
    _write_df(pm_ws, performance.monthly_returns)
    for row in range(2, len(performance.monthly_returns) + 2):
        for col in range(2, pm_ws.max_column + 1):
            if isinstance(pm_ws.cell(row, col).value, (float, int)):
                pm_ws.cell(row, col).number_format = '0.00%' if col in (5, 6) else '#,##0.00 €'
    if len(performance.monthly_returns) > 1:
        chart = LineChart()
        chart.title = "Performance cumulée estimée"
        data = Reference(pm_ws, min_col=6, min_row=1, max_row=len(performance.monthly_returns) + 1)
        cats = Reference(pm_ws, min_col=1, min_row=2, max_row=len(performance.monthly_returns) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        pm_ws.add_chart(chart, "H2")
    _format_sheet(pm_ws)

    ann_ws = wb.create_sheet("Rendements annuels")
    _write_df(ann_ws, performance.annual_returns)
    for row in range(2, len(performance.annual_returns) + 2):
        ann_ws.cell(row, 2).number_format = '0.00%'
        ann_ws.cell(row, 3).number_format = '#,##0.00 €'
        ann_ws.cell(row, 4).number_format = '#,##0.00 €'
    _format_sheet(ann_ws)

    attr_ws = wb.create_sheet("Attribution performance")
    _write_df(attr_ws, performance.attribution.head(50))
    for row in range(2, min(len(performance.attribution), 50) + 2):
        attr_ws.cell(row, 3).number_format = '#,##0.00 €'
        attr_ws.cell(row, 4).number_format = '0.00%'
    if len(performance.attribution) > 1:
        _add_bar(attr_ws, "Top contributions", 1, 3, 1, min(16, len(performance.attribution) + 1), "F2")
    _format_sheet(attr_ws)

    bench_ws = wb.create_sheet("Benchmark")
    _write_df(bench_ws, performance.benchmark)
    for row in range(2, len(performance.benchmark) + 2):
        for col in range(2, bench_ws.max_column + 1):
            if isinstance(bench_ws.cell(row, col).value, (float, int)):
                bench_ws.cell(row, col).number_format = '0.00%' if col == 4 else '#,##0.00'
    if len(performance.benchmark) > 1:
        chart = LineChart()
        chart.title = "Indice portefeuille vs benchmark 7%"
        data = Reference(bench_ws, min_col=2, max_col=3, min_row=1, max_row=len(performance.benchmark) + 1)
        cats = Reference(bench_ws, min_col=1, min_row=2, max_row=len(performance.benchmark) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        bench_ws.add_chart(chart, "F2")
    _format_sheet(bench_ws)



    # Look-through Dashboard v2.5
    lt_ws = wb.create_sheet("Look-through Dashboard")
    lt_ws["A1"] = "Look-through Dashboard v2.5"
    lt_ws["A1"].font = TITLE_FONT
    lt_ws["A2"] = "Exposition consolidée estimée aux sociétés sous-jacentes, secteurs, zones, devises et thèmes."
    lt_ws["A2"].font = SUBTITLE_FONT
    _write_df(lt_ws, lookthrough.company_exposure.head(25), start_row=4)
    for row in range(5, 5 + min(len(lookthrough.company_exposure), 25)):
        lt_ws.cell(row, 2).number_format = '#,##0.00 €'
        lt_ws.cell(row, 3).number_format = '0.00%'
    if len(lookthrough.company_exposure) > 0:
        _add_bar(lt_ws, "Top sociétés consolidées", 1, 2, 4, 4 + min(len(lookthrough.company_exposure), 15), "E4")

    lt_ws["A34"] = "Concentrations cachées"
    lt_ws["A34"].font = TITLE_FONT
    _write_df(lt_ws, lookthrough.hidden_concentration.head(25), start_row=36)
    for row in range(37, 37 + min(len(lookthrough.hidden_concentration), 25)):
        lt_ws.cell(row, 3).number_format = '#,##0.00 €'
        lt_ws.cell(row, 4).number_format = '0.00%'
    _format_sheet(lt_ws)

    lt_sector_ws = wb.create_sheet("LT Secteurs")
    _write_df(lt_sector_ws, lookthrough.sector_exposure)
    for row in range(2, len(lookthrough.sector_exposure) + 2):
        lt_sector_ws.cell(row, 2).number_format = '#,##0.00 €'
        lt_sector_ws.cell(row, 3).number_format = '0.00%'
    if len(lookthrough.sector_exposure) > 0:
        _add_pie(lt_sector_ws, "Secteurs réels", 1, 2, 1, len(lookthrough.sector_exposure) + 1, "E2")
    _format_sheet(lt_sector_ws)

    lt_geo_ws = wb.create_sheet("LT Géographie")
    _write_df(lt_geo_ws, lookthrough.geography_exposure)
    for row in range(2, len(lookthrough.geography_exposure) + 2):
        lt_geo_ws.cell(row, 2).number_format = '#,##0.00 €'
        lt_geo_ws.cell(row, 3).number_format = '0.00%'
    if len(lookthrough.geography_exposure) > 0:
        _add_pie(lt_geo_ws, "Zones réelles", 1, 2, 1, len(lookthrough.geography_exposure) + 1, "E2")
    _format_sheet(lt_geo_ws)

    lt_theme_ws = wb.create_sheet("LT Thèmes")
    _write_df(lt_theme_ws, lookthrough.theme_exposure)
    for row in range(2, len(lookthrough.theme_exposure) + 2):
        lt_theme_ws.cell(row, 2).number_format = '#,##0.00 €'
        lt_theme_ws.cell(row, 3).number_format = '0.00%'
    if len(lookthrough.theme_exposure) > 0:
        _add_bar(lt_theme_ws, "Thèmes réels", 1, 2, 1, len(lookthrough.theme_exposure) + 1, "E2")
    _format_sheet(lt_theme_ws)

    # Risk Dashboard v2.3 conservé
    risk = build_risk(transactions, net_worth, metrics)
    risk_ws = wb.create_sheet("Risk Dashboard")
    risk_ws["A1"] = "Risk Dashboard v2.3"
    risk_ws["A1"].font = TITLE_FONT
    risk_ws["A2"] = "Les métriques sont des estimations fondées sur des hypothèses de volatilité/corrélation par bloc d'actifs."
    risk_ws["A2"].font = SUBTITLE_FONT
    _write_df(risk_ws, risk.summary, start_row=4)
    for row in range(5, 5 + len(risk.summary)):
        label = str(risk_ws.cell(row, 1).value or "")
        if isinstance(risk_ws.cell(row, 2).value, (float, int)):
            if "VaR" in label or "CVaR" in label:
                risk_ws.cell(row, 2).number_format = '#,##0.00 €'
            elif "HHI" in label or "Nombre" in label or "Sharpe" in label:
                risk_ws.cell(row, 2).number_format = '#,##0.00'
            else:
                risk_ws.cell(row, 2).number_format = '0.00%'
    risk_ws["D4"] = "Alertes"
    risk_ws["D4"].font = TITLE_FONT
    _write_df(risk_ws, risk.alerts, start_row=6, start_col=4)
    risk_ws["A18"] = "Stress tests"
    risk_ws["A18"].font = TITLE_FONT
    _write_df(risk_ws, risk.stress_tests, start_row=20)
    for row in range(21, 21 + len(risk.stress_tests)):
        risk_ws.cell(row, 2).number_format = '0.00%'
        risk_ws.cell(row, 3).number_format = '#,##0.00 €'
    if len(risk.stress_tests) > 0:
        _add_bar(risk_ws, "Impact des stress tests", 1, 3, 20, 20 + len(risk.stress_tests), "E20")
    _format_sheet(risk_ws)

    rc_ws = wb.create_sheet("Contribution risque")
    _write_df(rc_ws, risk.risk_by_position)
    for row in range(2, len(risk.risk_by_position) + 2):
        for col in range(5, rc_ws.max_column + 1):
            if isinstance(rc_ws.cell(row, col).value, (float, int)):
                if col in (6, 7, 8):
                    rc_ws.cell(row, col).number_format = '0.00%'
                else:
                    rc_ws.cell(row, col).number_format = '#,##0.00'
    if len(risk.risk_by_position) > 1:
        _add_bar(rc_ws, "Top contributions au risque", 1, 8, 1, min(16, len(risk.risk_by_position) + 1), "J2")
    _format_sheet(rc_ws)

    corr_ws = wb.create_sheet("Corrélations estimées")
    if not risk.correlation_matrix.empty:
        corr_ws.append([""] + list(risk.correlation_matrix.columns))
        for idx, row in risk.correlation_matrix.iterrows():
            corr_ws.append([idx] + list(row.values))
        for cell in corr_ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for row in corr_ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                if isinstance(cell.value, (float, int)):
                    cell.number_format = '0.00'
    _format_sheet(corr_ws)

    meta = wb.create_sheet("Méthode")
    meta.append(["Point", "Détail"])
    for c in meta[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    meta.append(["Version", "2.6.0 Reports Dashboard"])
    meta.append(["Limite", "Les allocations sont estimées à partir des transactions ; elles ne remplacent pas une valorisation ligne par ligne."])
    meta.append(["PDF", f"Total détecté : {net_worth.total_value}"])
    meta.append(["Nouveauté v2.6", "Reports Dashboard : génération d’un rapport PDF professionnel en complément de l’export Excel enrichi. Le PDF synthétise les KPI, allocations, performance, risque, look-through, stress tests et méthodologie."])
    _format_sheet(meta)

    wb.save(output)
    return output
